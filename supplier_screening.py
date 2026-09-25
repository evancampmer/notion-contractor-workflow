import os
import io
import re
import json
import logging
import tempfile
import threading
from datetime import datetime

import requests
from dotenv import load_dotenv
from slack_bolt import App
from slack_bolt.adapter.flask import SlackRequestHandler
from flask import Flask, request
from openai import OpenAI
from docx import Document
from docx.shared import Pt, RGBColor
from google.oauth2 import service_account
from googleapiclient.discovery import build
from googleapiclient.http import MediaIoBaseDownload

# =========================
# CONFIG + LOGGING
# =========================

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s - %(levelname)s - %(message)s"
)

load_dotenv()

# =========================
# CLIENTS
# =========================

slack_app = App(
    token=os.getenv("SLACK_BOT_TOKEN"),
    signing_secret=os.getenv("SLACK_SIGNING_SECRET")
)

openrouter = OpenAI(
    base_url="https://openrouter.ai/api/v1",
    api_key=os.getenv("OPENROUTER_API_KEY")
)

MODEL = "anthropic/claude-sonnet-4-5"

SCOPES = ["https://www.googleapis.com/auth/drive"]
service_account_info = json.loads(os.getenv("GOOGLE_SERVICE_ACCOUNT_JSON"))
credentials = service_account.Credentials.from_service_account_info(
    service_account_info, scopes=SCOPES
)
drive_service = build("drive", "v3", credentials=credentials)

SLACK_CHANNEL_ID = "C0C2QRTGAUV"
GOOGLE_DRIVE_FOLDER_ID = os.getenv("GOOGLE_DRIVE_FOLDER_ID")
OUTPUT_FOLDER = "output_docs"
os.makedirs(OUTPUT_FOLDER, exist_ok=True)

# =========================
# PROMPTS
# =========================

CRITERIA_SYSTEM_PROMPT = """You are a business analyst at a consulting and staffing firm. Convert a project description into a focused screening rubric.

Output exactly two things with NO extra text:

1. A 1-2 sentence project summary.

2. A pipe-delimited criteria table:
| Category | Criterion | Importance | Evidence standard |

Rules:
- Category is "Technical" or "Contextual"
- 4-7 criteria total — only the most critical ones
- Use SHORT criterion names (2-4 words max, e.g. "GTM execution", "PMO cadence", "PE value creation")
- Importance: 1-5 (5 = must-have)
- Evidence standard: specific words/phrases to look for in a profile
- No commentary, no headers, no explanation outside these two items"""

BATCH_SCORING_PROMPT = """You are a supplier screening agent at a consulting and staffing firm.

Score EVERY supplier profile in this batch against the criteria. Return ONLY a pipe-delimited table — no other text, no preamble, no commentary.

Table columns (use exactly these headers):
| Name | Total Score /100 | Technical /100 | Contextual /100 | Each Criterion Score (0-5) | Strengths Summary | Weakness Summary |

Scoring rules:
- Score each criterion 0-5 based ONLY on explicit text evidence in the profile — no assumptions
- Multiply each score by its importance weight, sum all weighted scores, normalize to 100
- Technical subtotal and Contextual subtotal each normalized to 100 independently
- "Each Criterion Score": inline list using the short criterion names from the criteria table, e.g. "510k sub 5; FDA reg 4; Tech writing 2"
- Score ALL suppliers in this batch; give 0 if no relevant evidence exists
- Sort by Total Score descending
- Each cell: 1-2 sentences max"""

SYNTHESIS_PROMPT = """You are a supplier screening agent summarizing final candidate scores.

You will receive the top-scored candidates after all batches have been scored. Write exactly three sections:

SECTION 1 — HEADLINE
One sentence: overall finding (e.g. "Three strong candidates identified; no exact FDA regulatory specialist found in database")

SECTION 2 — SCORING METHOD
2-3 sentences: how technical vs contextual were weighted, which criteria drove the most differentiation, and the evidence standard applied.

SECTION 3 — RECOMMENDATION
One bullet per recommended candidate (top 5 maximum). Each bullet: name, action ("Advance as primary" or "Advance as alternative"), 1-2 sentences on fit and key gap. End with one sentence on any universal gap across all candidates.

No other text outside these three sections."""

# =========================
# GOOGLE DRIVE
# =========================

def get_supplier_docs():
    """Download all DOCX files from the Google Drive folder and extract text."""
    supplier_docs = []

    query = (
        f"'{GOOGLE_DRIVE_FOLDER_ID}' in parents "
        f"and mimeType='application/vnd.openxmlformats-officedocument.wordprocessingml.document' "
        f"and trashed=false"
    )

    results = drive_service.files().list(
        q=query,
        fields="files(id, name)",
        supportsAllDrives=True,
        includeItemsFromAllDrives=True,
        corpora="allDrives"
    ).execute()

    files = results.get("files", [])
    logging.info(f"Found {len(files)} supplier docs in Drive")

    for file in files:
        try:
            req = drive_service.files().get_media(
                fileId=file["id"],
                supportsAllDrives=True
            )
            buffer = io.BytesIO()
            downloader = MediaIoBaseDownload(buffer, req)
            done = False
            while not done:
                _, done = downloader.next_chunk()

            buffer.seek(0)
            with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp:
                tmp.write(buffer.read())
                tmp_path = tmp.name

            doc = Document(tmp_path)
            text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())

            supplier_docs.append({
                "name": file["name"].replace(".docx", "").replace("_", " "),
                "text": text
            })

            os.unlink(tmp_path)

        except Exception as e:
            logging.warning(f"Failed to read {file['name']}: {e}")

    return supplier_docs


# =========================
# TEXT HELPERS
# =========================

def strip_mention(text):
    “””Remove Slack @mention tags like <@U0C407F3AKT> from the message.”””
    return re.sub(r”<@[A-Z0-9]+>\s*”, “”, text).strip()


def _strip_md(text):
    “””Strip markdown bold/italic markers.”””
    text = re.sub(r”\*\*(.+?)\*\*”, r”\1”, text)
    text = re.sub(r”\*(.+?)\*”, r”\1”, text)
    return text.strip()


def _parse_md_table(text):
    “””Extract rows from a markdown pipe-delimited table.”””
    rows = []
    for line in text.split(“\n”):
        line = line.strip()
        if not line.startswith(“|”):
            continue
        if re.match(r”^\|[-|:\s]+\|$”, line):
            continue  # separator row
        cells = [_strip_md(c.strip()) for c in line.strip(“|”).split(“|”)]
        if any(cells):
            rows.append(cells)
    return rows


# =========================
# AI CALLS
# =========================

def create_supplier_criteria(message_text):
    logging.info("Creating supplier criteria...")
    response = openrouter.chat.completions.create(
        model=MODEL,
        messages=[
            {"role": "system", "content": CRITERIA_SYSTEM_PROMPT},
            {"role": "user", "content": (
                "Create the screening criteria for this project brief:\n\n"
                f"{message_text}"
            )}
        ]
    )
    return response.choices[0].message.content


BATCH_SIZE = 30

def _score_batch(criteria, batch_docs):
    """Score one batch of supplier docs and return the parsed table rows."""
    supplier_text = "\n\n---\n\n".join(
        f"SUPPLIER: {doc['name']}\n\n{doc['text']}"
        for doc in batch_docs
    )
    response = openrouter.chat.completions.create(
        model=MODEL,
        messages=[
            {"role": "system", "content": BATCH_SCORING_PROMPT},
            {"role": "user", "content": f"CRITERIA:\n{criteria}\n\nSUPPLIER PROFILES:\n{supplier_text}"}
        ]
    )
    return response.choices[0].message.content


def _extract_score(row):
    """Parse the Total Score from a table row (column index 1)."""
    try:
        return int(re.search(r"\d+", row[1]).group())
    except Exception:
        return 0


def _is_header_row(row):
    """Detect if a row is a header (not a data row)."""
    return any(h in row[0].lower() for h in ["name", "rank", "supplier"])


def rank_suppliers(criteria, supplier_docs):
    """Score all suppliers in batches, merge results, return top candidates + synthesis."""
    batches = [supplier_docs[i:i+BATCH_SIZE] for i in range(0, len(supplier_docs), BATCH_SIZE)]
    logging.info(f"Scoring {len(supplier_docs)} suppliers across {len(batches)} batches of up to {BATCH_SIZE}")

    header_row = None
    all_data_rows = []

    for idx, batch in enumerate(batches):
        logging.info(f"Batch {idx+1}/{len(batches)}: scoring {len(batch)} suppliers...")
        batch_text = _score_batch(criteria, batch)
        rows = _parse_md_table(batch_text)

        if not rows:
            logging.warning(f"Batch {idx+1} returned no parseable rows")
            continue

        for row in rows:
            if _is_header_row(row):
                if header_row is None:
                    header_row = row
            else:
                all_data_rows.append(row)

    # Sort all candidates by Total Score, take top 15
    all_data_rows.sort(key=_extract_score, reverse=True)
    top_rows = all_data_rows[:15]
    logging.info(f"Merged {len(all_data_rows)} scored candidates, kept top {len(top_rows)}")

    # Build synthesis (headline + scoring method + recommendation) from top results
    top_summary = "\n".join(
        f"{row[0]}: Total {row[1]}, Tech {row[2] if len(row) > 2 else 'N/A'}, "
        f"Contextual {row[3] if len(row) > 3 else 'N/A'}, "
        f"Criteria scores: {row[4] if len(row) > 4 else 'N/A'}, "
        f"Strengths: {row[5] if len(row) > 5 else 'N/A'}"
        for row in top_rows
    )

    synthesis_response = openrouter.chat.completions.create(
        model=MODEL,
        messages=[
            {"role": "system", "content": SYNTHESIS_PROMPT},
            {"role": "user", "content": f"CRITERIA:\n{criteria}\n\nTOP CANDIDATES (pre-scored):\n{top_summary}"}
        ]
    )
    synthesis = synthesis_response.choices[0].message.content
    logging.info("Synthesis complete")

    return {
        "synthesis": synthesis,
        "header_row": header_row,
        "ranked_rows": top_rows
    }


# =========================
# DOCX GENERATION
# =========================

def _add_table_to_doc(doc, rows):
    """Add a parsed markdown table to a python-docx Document."""
    if not rows:
        return
    num_cols = len(rows[0])
    table = doc.add_table(rows=len(rows), cols=num_cols)
    table.style = "Table Grid"
    for r_idx, row in enumerate(rows):
        # Pad/trim to expected column count
        while len(row) < num_cols:
            row.append("")
        row = row[:num_cols]
        for c_idx, cell_text in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.text = cell_text
            para = cell.paragraphs[0]
            if r_idx == 0:
                for run in para.runs:
                    run.bold = True
    doc.add_paragraph()


def _parse_rankings(rankings_text):
    """Split the AI screening output into its four sections.
    Handles varied formatting: plain, bold (**SECTION 1**), markdown headers (## SECTION 1), etc.
    """
    sections = {1: [], 2: [], 3: [], 4: []}
    current = None
    for line in rankings_text.split("\n"):
        # Strip markdown bold/header markers before checking for section labels
        clean = re.sub(r"[#\*\_]+", "", line).strip()
        if re.search(r"SECTION\s*1", clean, re.IGNORECASE):
            current = 1; continue
        if re.search(r"SECTION\s*2", clean, re.IGNORECASE):
            current = 2; continue
        if re.search(r"SECTION\s*3", clean, re.IGNORECASE):
            current = 3; continue
        if re.search(r"SECTION\s*4", clean, re.IGNORECASE):
            current = 4; continue
        if current is not None:
            sections[current].append(line)
    return {k: "\n".join(v).strip() for k, v in sections.items()}


def generate_recommendations_docx(criteria, rankings, original_message):
    """Build the output DOCX matching Cassidy's style.

    `rankings` is a dict returned by rank_suppliers():
        {
            "synthesis": <str — the 3-section AI synthesis>,
            "header_row": <list — column headers from batch scoring>,
            "ranked_rows": <list of lists — top-15 candidate rows>,
        }
    """
    doc = Document()

    # Derive a short title from the first non-empty line of the message
    title_line = next((l.strip() for l in original_message.split("\n") if l.strip()), "Supplier Screening")
    if len(title_line) > 80:
        title_line = title_line[:77] + "..."

    # Main heading — matches Cassidy's style exactly
    doc.add_heading(f"Supplier screening results — {title_line}", level=2)

    # ── Scoring method ──────────────────────────────────────────
    doc.add_heading("Scoring method", level=3)

    # Criteria summary (prose before the pipe table)
    criteria_lines = criteria.split("\n")
    summary_parts, table_lines = [], []
    in_table = False
    for line in criteria_lines:
        if "|" in line:
            in_table = True
        (table_lines if in_table else summary_parts).append(line)

    summary_text = " ".join(l.strip() for l in summary_parts if l.strip())
    if summary_text:
        doc.add_paragraph(summary_text)

    # Criteria table
    criteria_rows = _parse_md_table("\n".join(table_lines))
    if criteria_rows:
        _add_table_to_doc(doc, criteria_rows)

    # Parse synthesis into sections 1-3
    synthesis_text = rankings.get("synthesis", "")
    sections = _parse_rankings(synthesis_text)
    logging.info(f"Parsed synthesis sections — lengths: { {k: len(v) for k, v in sections.items()} }")

    # Section 1 — headline (brief overall finding paragraph)
    if sections[1]:
        doc.add_paragraph(_strip_md(sections[1]))

    # Section 2 — scoring method prose
    if sections[2]:
        doc.add_paragraph(_strip_md(sections[2]))

    # ── Recommendation ──────────────────────────────────────────
    doc.add_heading("Recommendation", level=3)
    for line in sections[3].split("\n"):
        s = line.strip()
        if not s:
            continue
        if re.match(r"^[\*\-]\s", s):
            content = _strip_md(re.sub(r"^[\*\-]\s+", "", s))
            doc.add_paragraph(content, style="List Paragraph")
        else:
            doc.add_paragraph(_strip_md(s))

    # ── Ranked candidate table (from merged batch results) ───────
    header_row = rankings.get("header_row")
    ranked_rows = rankings.get("ranked_rows", [])
    if header_row and ranked_rows:
        _add_table_to_doc(doc, [header_row] + ranked_rows)
    elif ranked_rows:
        _add_table_to_doc(doc, ranked_rows)

    filename = os.path.join(OUTPUT_FOLDER, "Supplier_Recommendations.docx")
    doc.save(filename)
    return filename


# =========================
# PIPELINE
# =========================

def process_message(message_text, channel_id, thread_ts, client):
    try:
        criteria = create_supplier_criteria(message_text)
        logging.info("Criteria created")

        supplier_docs = get_supplier_docs()

        if not supplier_docs:
            client.chat_postMessage(
                channel=channel_id, thread_ts=thread_ts,
                text="No supplier profiles found in the Google Drive folder."
            )
            return

        rankings = rank_suppliers(criteria, supplier_docs)
        logging.info("Rankings complete")

        docx_path = generate_recommendations_docx(criteria, rankings, message_text)

        client.files_upload_v2(
            channel=channel_id,
            thread_ts=thread_ts,
            file=docx_path,
            filename="Supplier_Recommendations.docx",
            initial_comment="Here are your supplier recommendations! 📋"
        )

        logging.info("Supplier recommendations sent to Slack")

    except Exception as e:
        logging.error(f"Pipeline failed: {e}", exc_info=True)
        client.chat_postMessage(
            channel=channel_id, thread_ts=thread_ts,
            text=f"Something went wrong: {str(e)}"
        )


# =========================
# SLACK EVENT HANDLER
# =========================

@slack_app.event("message")
def handle_message(event, client):
    if event.get("bot_id") or event.get("subtype"):
        return

    channel_id = event.get("channel")
    if channel_id != SLACK_CHANNEL_ID:
        return

    raw_text = event.get("text", "").strip()
    message_text = strip_mention(raw_text)

    if not message_text:
        return

    thread_ts = event.get("ts")
    logging.info(f"Received message: {message_text[:100]}...")

    client.chat_postMessage(
        channel=channel_id,
        thread_ts=thread_ts,
        text="Got it! Analyzing your request and scanning the supplier database... this may take a minute. 🔍"
    )

    thread = threading.Thread(
        target=process_message,
        args=(message_text, channel_id, thread_ts, client)
    )
    thread.daemon = True
    thread.start()


# =========================
# FLASK SERVER
# =========================

flask_app = Flask(__name__)
handler = SlackRequestHandler(slack_app)


@flask_app.route("/slack/events", methods=["POST"])
def slack_events():
    return handler.handle(request)


@flask_app.route("/health", methods=["GET"])
def health():
    return "OK", 200


if __name__ == "__main__":
    flask_app.run(port=int(os.getenv("PORT", 3000)))